"""Excelを唯一の正(SSOT)にするスケジュール管理ブックを作る。

使い方:
    python -m tools.make_schedule_workbook data/workshop_schedule.xlsx            # 空のブック(タスクマスタのみ)
    python -m tools.make_schedule_workbook data/sample_schedule.xlsx --sample   # サンプルデータ入り

※ openpyxl で作ったファイルには数式の計算結果が入っていない。Excel(ブラウザ版含む)で
  一度開けば自動で計算される(fullCalcOnLoad)。Webアプリは計算結果を読むので、
  作成後は一度 Excel で開いて保存してから使うこと。

ブックの考え方(VBA/Office Scripts を使わず、ブラウザ版Excel・個人用OneDriveで動く関数だけを使う):
  - タスク一覧は「1行 = 1タスク」の入力表。WSのタスクは、WSを作るときに「タスク生成」シートで
    タスクマスタから作った行を「値のみ貼り付け」して作る。値として固定されるので、後でマスタを
    変えても既存のWSのタスクは変わらない(変更は今後生成するWSにだけ効く)。
  - WS固有のタスクは、タスク一覧に行を足すだけ。
  - 開始日・終了日・作業日数・警告などは、タスク一覧の数式の列で自動計算する
    (日付決定有無=Yes のタスクの日付を入れると、後続タスクの日付が決まる)。
  - 日付の計算は「同じWSで並び順が小さいタスクが上の行にある」前提(循環参照を避けるため)。
    崩れていると警告列に出るので、WS名→並び順で並べ替える。
"""

import argparse
from datetime import datetime

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

_parser = argparse.ArgumentParser(description="スケジュール管理ブック(Excel)を作る")
_parser.add_argument("output", nargs="?", default="data/workshop_schedule.xlsx")
_parser.add_argument("--sample", action="store_true", help="サンプルのWS・メンバー・タスクを入れる")
_args = _parser.parse_args()
OUT = _args.output
SAMPLE = _args.sample

FONT = "Arial"
N_WS = 100  # WSシートの入力行数
N_MASTER = 50  # タスクマスタの入力行数
N_TASK = 1000  # タスク一覧に数式を入れる行数
N_OTHER = 200
N_MEMBER = 50
N_ASSIGNEES = 10  # 1タスクの担当者の最大人数
STATUSES = ("未着手", "着手中", "完了")
GEN_FIRST = 8  # タスク生成シートの表の最初の行

HDR_INPUT = PatternFill("solid", start_color="305496")  # 入力するシート/列の見出し
HDR_AUTO = PatternFill("solid", start_color="548235")  # 自動計算の見出し
HDR_HELP = PatternFill("solid", start_color="808080")  # 計算用
INPUT_FILL = PatternFill("solid", start_color="FFF9E5")
AUTO_FILL = PatternFill("solid", start_color="F2F2F2")
NO_FILL = PatternFill()
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
DATE = "yyyy/mm/dd"
STATUS_FILLS = {"未着手": "FFF2CC", "着手中": "DDEBF7", "完了": "D9D9D9"}
L = get_column_letter

wb = Workbook()


def base_font(bold=False, color="000000", size=10, italic=False):
    return Font(name=FONT, bold=bold, color=color, size=size, italic=italic)


def header(ws, cols, fills, notes=None, row=1):
    for i, (name, width) in enumerate(cols, start=1):
        c = ws.cell(row=row, column=i, value=name or None)
        c.font = base_font(bold=True, color="FFFFFF")
        c.fill = fills[i - 1] if isinstance(fills, list) else fills
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        if name:
            c.border = BORDER
        ws.column_dimensions[L(i)].width = width
        if notes and notes.get(name):
            c.comment = Comment(notes[name], "template")
    ws.row_dimensions[row].height = 30
    ws.freeze_panes = f"A{row + 1}"


def body_style(ws, first_col, last_col, first_row, last_row, fill, date_cols=(), font_color="000000"):
    for r in range(first_row, last_row + 1):
        for col in range(first_col, last_col + 1):
            c = ws.cell(row=r, column=col)
            c.font = base_font(color=font_color)
            c.fill = fill
            c.border = BORDER
            if L(col) in date_cols:
                c.number_format = DATE


def add_list(ws, rng, formula):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "入力エラー"
    dv.error = "一覧から選んでください"
    ws.add_data_validation(dv)
    dv.add(rng)


def add_date_rule(ws, rng):
    dv = DataValidation(type="date", operator="greaterThan", formula1="1", allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "入力エラー"
    dv.error = "日付を入力してください(例: 2026/10/01)"
    ws.add_data_validation(dv)
    dv.add(rng)


def add_int_rule(ws, rng):
    dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="1", allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = "入力エラー"
    dv.error = "1以上の整数を入力してください"
    ws.add_data_validation(dv)
    dv.add(rng)


def d(s):
    return datetime.strptime(s, "%Y/%m/%d")


WS_LAST = N_WS + 1
M_LAST = N_MASTER + 1
T_LAST = N_TASK + 1
TQ = "'タスク一覧'"
TMQ = "'タスクマスタ'"

# ---------------------------------------------------------------------------
# 説明
# ---------------------------------------------------------------------------
ws = wb.active
ws.title = "説明"
ws.sheet_view.showGridLines = False
ws.column_dimensions["A"].width = 2
ws.column_dimensions["B"].width = 22
ws.column_dimensions["C"].width = 110
ws["B2"] = "ワークショップ スケジュール管理ブック(このファイルが唯一の正)"
ws["B2"].font = base_font(bold=True, size=14)
lines = [
    ("このファイルの役割", ""),
    ("", "スケジュールの正式なデータはこのファイルだけです。Webアプリはこのファイルを読み取って表示するだけで、書き換えません。"),
    ("", "マクロは使っていません(ブラウザ版Excelでそのまま動きます)。日付の計算は数式で自動的に行われます。"),
    ("", ""),
    ("シートの見分け方", ""),
    ("青い見出し", "入力する列です。薄い黄色の範囲に記入します。"),
    ("緑の見出し", "自動計算の列です。直接入力しないでください(数式が消えます)。"),
    ("灰色の見出し", "計算用の列です(折りたたんで非表示にしています)。変更しないでください。"),
    ("", ""),
    ("新しいWSを追加する", ""),
    ("1", "「WS」シートの空いている行に WS名・種別(ワークショップ)・並び順・開始日 を入力します。「タスク」列が「未生成」になります。"),
    ("2", "「タスク生成」シートの B1 でそのWSを選ぶと、タスクマスタ(WS自動生成=対象)のタスクが表に並びます。"),
    ("3", "表の WS名〜日付決定有無(A〜E列)の行を選んでコピーし、「タスク一覧」の一番下の空き行の A列に「値のみ貼り付け」します。"),
    ("", "必ず「値のみ」で貼り付けてください(普通に貼ると数式のまま入り、警告列に表示されます)。"),
    ("", "値として貼り付けるので、後でタスクマスタを変えてもこのWSのタスクは変わりません。"),
    ("4", "WS固有のタスクは、タスク一覧の一番下の空き行に WS名・タスク名・並び順・所要日数・日付決定有無 を入力して追加します。"),
    ("5", "並び順が前後したら、タスク一覧を WS名 → 並び順 で並べ替えます(データ → 並べ替え)。"),
    ("", ""),
    ("日付を決める・担当者を入れる", ""),
    ("1", f"「タスク一覧」の各行の 開始日(手入力)・終了日(手入力)・ステータス・担当者1〜{N_ASSIGNEES}・メモ に入力します。"),
    ("2", "Yes のタスク(リハーサル・WS実施)に開始日(手入力)を入れると、その後のタスクの日付が自動で入ります。"),
    ("3", "No のタスクに開始日(手入力)を入れると「手動」になり、その日付で固定されます(空欄なら自動計算)。"),
    ("4", "終了日(手入力)を入れると、所要日数ではなくその日付で終わります。"),
    ("5", "所要日数を変えると、そのタスク(とその後のタスク)の日付だけが計算し直されます。"),
    ("ステータス", "未着手 / 着手中 / 完了 から選びます(空欄は未着手)。Webアプリではバーの色で表示されます。"),
    ("絞り込み", "タスク一覧・個別タスクの見出しのフィルター(▼)で、WS名やステータスで絞り込めます。"),
    ("", ""),
    ("日付の計算ルール", ""),
    ("並び順1", "WS の開始日から開始。WS の開始日が空欄だと、そのWSのタスクは「未定」になります。"),
    ("Yes の次", "日付決定有無=Yes のタスクの次の並び順は、その Yes タスクの終了日の翌日から開始。"),
    ("それ以外", "直前の並び順のタスクのうち最も遅い終了日の翌日から開始(直列)。同じ並び順のタスクは同じ期間(並列)。"),
    ("並び順が空欄", "WS固有の単発作業など。開始日(手入力)に入れた日付で表示します(空欄なら未定)。"),
    ("未定", "起点になるタスクに1件でも未定があれば、その後の自動のタスクも未定。"),
    ("終了日", "終了日(手入力)があればその日付、無ければ 開始日 + 所要日数 − 1(暦日)。"),
    ("", ""),
    ("その他の作業", ""),
    ("個別タスク", "社内MTGなど、WSに属さない作業は「個別タスク」シートに直接入力します。WS名は空欄でよく、画面では「その他」の行に出ます。"),
    ("", ""),
    ("注意", ""),
    ("並び順と行の順", "日付の計算は「同じWSで並び順が小さいタスクが上の行にある」ことが前提です。"),
    ("", "崩れていると警告列に「並び順が小さいタスクが下の行にあります」と出るので、WS名 → 並び順で並べ替えてください。"),
    ("行の追加", "タスク一覧の行は一番下の空き行に追加してください。途中に行を挿入すると、その行には計算の数式が入りません。"),
    ("", "行の削除・並べ替えは自由にできます。タスクマスタ・WS の行も自由に編集・並べ替えできます。"),
    ("タスクマスタ", "マスタの変更は、これから生成するWSにだけ反映されます(生成済みのWSのタスクは変わりません)。"),
    ("", "生成済みのWSも直したいときは、タスク一覧のそのWSの行を直接直してください。"),
    ("WS名の変更", "WS名を変えるときは、タスク一覧・個別タスクの WS名 も同じ名前に直してください。"),
]
r = 4
for label, text in lines:
    b, c = ws.cell(row=r, column=2, value=label or None), ws.cell(row=r, column=3, value=text or None)
    heading = text == "" and label != ""
    b.font = base_font(bold=True, size=12 if heading else 10, color="305496" if heading else "000000")
    c.font = base_font()
    c.alignment = Alignment(wrap_text=True, vertical="top")
    b.alignment = Alignment(vertical="top")
    r += 1
for cell, fill in (("B9", HDR_INPUT), ("B10", HDR_AUTO), ("B11", HDR_HELP)):
    ws[cell].fill = fill
    ws[cell].font = base_font(bold=True, color="FFFFFF")

# ---------------------------------------------------------------------------
# WS
# ---------------------------------------------------------------------------
ws_ws = wb.create_sheet("WS")
header(
    ws_ws,
    [("WS名", 22), ("種別", 14), ("並び順", 9), ("開始日", 12), ("ステータス", 12), ("タスク\n(自動)", 12)],
    [HDR_INPUT] * 5 + [HDR_AUTO],
    {
        "WS名": "重複不可。タスク一覧・個別タスクはこの名前で紐づけます",
        "種別": "ワークショップ = タスク生成シートでタスクを作る / その他 = 個別タスク用(この行が無くても個別タスクは「その他」に出ます)",
        "開始日": "並び順1のタスクの開始日になります",
    },
)
body_style(ws_ws, 1, 5, 2, WS_LAST, INPUT_FILL, date_cols=("D",))
body_style(ws_ws, 6, 6, 2, WS_LAST, AUTO_FILL)
for r in range(2, WS_LAST + 1):
    ws_ws[f"F{r}"] = (
        f'=IF($A{r}="","",IF($B{r}="その他","―",IF(COUNTIF({TQ}!$A$2:$A${T_LAST},$A{r})=0,"未生成",'
        f'COUNTIF({TQ}!$A$2:$A${T_LAST},$A{r})&"件")))'
    )
    ws_ws[f"F{r}"].alignment = Alignment(horizontal="center")
SAMPLE_WS = [
    ["WS 秋季研修", "ワークショップ", 1, d("2026/10/01"), "準備中"],
    ["Snow Dorm", "ワークショップ", 2, d("2026/11/02"), "企画中"],
    ["その他", "その他", 99, None, None],
]
for i, row in enumerate(SAMPLE_WS if SAMPLE else [], start=2):
    for j, v in enumerate(row, start=1):
        ws_ws.cell(row=i, column=j, value=v)
add_list(ws_ws, f"B2:B{WS_LAST}", '"ワークショップ,その他"')
add_date_rule(ws_ws, f"D2:D{WS_LAST}")
add_int_rule(ws_ws, f"C2:C{WS_LAST}")
ws_ws.conditional_formatting.add(
    f"F2:F{WS_LAST}",
    FormulaRule(formula=['$F2="未生成"'], fill=PatternFill("solid", start_color="F8CBAD"),
                font=Font(name=FONT, color="C00000", bold=True)),
)

# ---------------------------------------------------------------------------
# タスクマスタ
# ---------------------------------------------------------------------------
ws_m = wb.create_sheet("タスクマスタ")
header(
    ws_m,
    [("タスク名", 22), ("標準作業日数", 12), ("標準並び順", 11), ("日付決定有無", 12), ("WS自動生成", 12),
     ("生成順\n(自動)", 10)],
    [HDR_INPUT] * 5 + [HDR_AUTO],
    {
        "標準並び順": "同じ値のタスクは並列(同じ期間)になります",
        "日付決定有無": "Yes = 人が日付を決めるタスク。決まるまで開始日は未定で、以降のタスクも未定になります",
        "WS自動生成": "空欄 = 対象。対象外にするとタスク生成シートに出ません",
    },
)
body_style(ws_m, 1, 5, 2, M_LAST, INPUT_FILL)
body_style(ws_m, 6, 6, 2, M_LAST, AUTO_FILL)
for r in range(2, M_LAST + 1):
    # 生成順 = 標準並び順の小さい順(同じ並び順なら上の行から)の通し番号
    ws_m[f"F{r}"] = (
        f'=IF(AND($A{r}<>"",$E{r}<>"対象外"),'
        f'COUNTIFS($A$2:$A${M_LAST},"<>",$E$2:$E${M_LAST},"<>対象外",$C$2:$C${M_LAST},"<"&$C{r})'
        f'+COUNTIFS($A$2:$A{r},"<>",$E$2:$E{r},"<>対象外",$C$2:$C{r},$C{r}),"")'
    )
MASTERS = [
    ("コンテンツ開発", 28, 1, "No"),
    ("場所決め", 28, 1, "No"),
    ("リハーサル物品購入", 28, 1, "No"),
    ("リハーサル&広報撮影", 1, 2, "Yes"),
    ("募集フォーム作成", 14, 3, "No"),
    ("広報動画作成", 14, 3, "No"),
    ("LP作成", 14, 3, "No"),
    ("広報稼働", 28, 4, "No"),
    ("物品購入", 28, 4, "No"),
    ("ワークショップ実施", 1, 5, "Yes"),
    ("アンケート分析", 7, 6, "No"),
]
for i, (name, days, order, yes) in enumerate(MASTERS, start=2):
    for j, v in enumerate([name, days, order, yes, "対象"], start=1):
        ws_m.cell(row=i, column=j, value=v)
add_list(ws_m, f"D2:D{M_LAST}", '"Yes,No"')
add_list(ws_m, f"E2:E{M_LAST}", '"対象,対象外"')
add_int_rule(ws_m, f"B2:C{M_LAST}")

# ---------------------------------------------------------------------------
# メンバー
# ---------------------------------------------------------------------------
ws_mem = wb.create_sheet("メンバー")
header(
    ws_mem,
    [("姓", 10), ("名", 10), ("表示名", 16), ("メールアドレス", 26), ("状態", 8)],
    HDR_INPUT,
    {"表示名": "重複不可。担当者のプルダウンに出る名前です"},
)
body_style(ws_mem, 1, 5, 2, N_MEMBER + 1, INPUT_FILL)
for i, row in enumerate(
    [] if not SAMPLE else [
        ["佐藤", "太郎", "佐藤 太郎", "sato@example.com", "有効"],
        ["鈴木", "次郎", "鈴木 次郎", "suzuki@example.com", "有効"],
        ["田中", "三郎", "田中 三郎", "tanaka@example.com", "有効"],
        ["村上", "四郎", "村上 四郎", "murakami@example.com", "有効"],
    ],
    start=2,
):
    for j, v in enumerate(row, start=1):
        ws_mem.cell(row=i, column=j, value=v)
add_list(ws_mem, f"E2:E{N_MEMBER + 1}", '"有効,無効"')
MEMBERS_LIST = f"=メンバー!$C$2:$C${N_MEMBER + 1}"

# ---------------------------------------------------------------------------
# タスク一覧(1行 = 1タスク。左の5列と手入力の列に入力し、緑の列は自動計算)
# ---------------------------------------------------------------------------
# 列: A WS名 / B タスク名 / C 並び順 / D 所要日数 / E 日付決定有無   … タスク生成から値で貼り付け(または手入力)
#     F 開始日 / G 終了日 / H 作業日数 / I 担当者 / J 開始日の決め方 / K 警告   … 自動計算
#     L 開始日(手入力) / M 終了日(手入力) / N ステータス / O〜X 担当者1〜10 / Y メモ   … 入力
#     AA WS開始日 / AB 前の並び順 / AC 起点日 / AD WS順   … 計算用(非表示)
ws_t = wb.create_sheet("タスク一覧", 1)
AS1, ASN = L(15), L(14 + N_ASSIGNEES)  # O, X
MEMO = L(15 + N_ASSIGNEES)  # Y
cols = (
    [("WS名", 18), ("タスク名", 22), ("並び順", 7), ("所要日数", 8), ("日付決定有無", 10),
     ("開始日", 11), ("終了日", 11), ("作業日数", 8), ("担当者", 24), ("開始日の決め方", 12), ("警告", 36),
     ("開始日(手入力)", 12), ("終了日(手入力)", 12), ("ステータス", 9)]
    + [(f"担当者{i}", 11) for i in range(1, N_ASSIGNEES + 1)]
    + [("メモ", 24), ("", 2), ("WS開始日", 11), ("前の並び順", 9), ("起点日", 11), ("WS順", 6)]
)
fills = [HDR_INPUT] * 5 + [HDR_AUTO] * 6 + [HDR_INPUT] * (4 + N_ASSIGNEES) + [NO_FILL] + [HDR_HELP] * 4
header(
    ws_t,
    cols,
    fills,
    {
        "WS名": "タスク生成シートから値のみ貼り付け、または手入力(WSシートにある名前)",
        "並び順": "小さい順に実行。同じ値は並列。空欄 = 単発の作業(開始日(手入力)の日付で表示)",
        "所要日数": "作業にかかる日数(暦日)。終了日 = 開始日 + 所要日数 − 1",
        "日付決定有無": "Yes = 人が日付を決めるタスク。開始日(手入力)を入れるまで未定",
        "開始日": "自動計算。直接入力しないで「開始日(手入力)」に入れてください",
        "開始日の決め方": "決定待ち / 決定 = Yesのタスク。手動 = 開始日(手入力)で指定。自動 = ルールで計算",
        "警告": "並び順の上下・WSの開始日・日付の前後関係・貼り付け方の問題を表示",
        "開始日(手入力)": "Yesのタスク: 決定した日付 / Noのタスク: 入れると手動で固定、空欄なら自動計算",
        "終了日(手入力)": "入れるとその日付で終了(空欄なら 開始日+所要日数−1)",
        "ステータス": "未着手 / 着手中 / 完了(空欄は未着手)",
    },
)
ws_t.freeze_panes = "C2"
body_style(ws_t, 1, 5, 2, T_LAST, INPUT_FILL)
body_style(ws_t, 6, 11, 2, T_LAST, AUTO_FILL, date_cols=("F", "G"))
body_style(ws_t, 12, 15 + N_ASSIGNEES, 2, T_LAST, INPUT_FILL, date_cols=("L", "M"))
body_style(ws_t, 27, 30, 2, T_LAST, AUTO_FILL, date_cols=("AA", "AC"), font_color="595959")

WS_A = f"WS!$A$1:$A${WS_LAST}"
WS_D = f"WS!$D$1:$D${WS_LAST}"
for r in range(2, T_LAST + 1):
    p = r - 1  # 直前の行(範囲は見出し行1から直前の行まで=自分の行を含めない → 循環参照にならない)
    up = lambda col: f"${col}$1:${col}{p}"  # noqa: E731
    f = {}
    # --- 計算用 ---
    f["AA"] = (
        f'=IF($A{r}="","",IFERROR(IF(INDEX({WS_D},MATCH($A{r},{WS_A},0))="","未定",'
        f'INDEX({WS_D},MATCH($A{r},{WS_A},0))),"未定"))'
    )
    f["AB"] = (
        f'=IF(OR($A{r}="",$C{r}=""),"",IF(COUNTIFS({up("A")},$A{r},{up("C")},"<"&$C{r})=0,"なし",'
        f'_xlfn.MAXIFS({up("C")},{up("A")},$A{r},{up("C")},"<"&$C{r})))'
    )
    same_prev = f'{up("A")},$A{r},{up("C")},$AB{r}'
    f["AC"] = (
        f'=IF(OR($A{r}="",$C{r}=""),"",IF($AB{r}="なし",$AA{r},'
        f'IF(COUNTIFS({same_prev},{up("E")},"Yes")>0,'
        f'IF(COUNTIFS({same_prev},{up("E")},"Yes",{up("F")},"未定")>0,"未定",'
        f'_xlfn.MAXIFS({up("G")},{same_prev},{up("E")},"Yes")+1),'
        f'IF(COUNTIFS({same_prev},{up("F")},"未定")>0,"未定",'
        f'_xlfn.MAXIFS({up("G")},{same_prev})+1))))'
    )
    f["AD"] = f'=IF($A{r}="","",IFERROR(MATCH($A{r},{WS_A},0),""))'
    # --- 表示列 ---
    f["F"] = (
        f'=IF($A{r}="","",IF($C{r}="",IF($L{r}<>"",$L{r},"未定"),'
        f'IF($E{r}="Yes",IF($L{r}="","未定",$L{r}),IF($L{r}<>"",$L{r},$AC{r}))))'
    )
    f["G"] = f'=IF($A{r}="","",IF($M{r}<>"",$M{r},IF(ISNUMBER($F{r}),$F{r}+MAX(N($D{r}),1)-1,"未定")))'
    f["H"] = f'=IF($A{r}="","",IF(AND(ISNUMBER($F{r}),ISNUMBER($G{r})),$G{r}-$F{r}+1,MAX(N($D{r}),1)))'
    f["I"] = f'=IF($A{r}="","",_xlfn.TEXTJOIN("、",TRUE,${AS1}{r}:${ASN}{r}))'
    f["J"] = (
        f'=IF($A{r}="","",IF($C{r}="","手動",IF($E{r}="Yes",IF($L{r}="","決定待ち","決定"),'
        f'IF($L{r}<>"","手動",IF($M{r}<>"","自動(終了日指定)","自動")))))'
    )
    order_check = (
        f'AND($C{r}<>"",COUNTIFS($A{r + 1}:$A${T_LAST},$A{r},$C{r + 1}:$C${T_LAST},"<"&$C{r})>0)'
        if r < T_LAST else "FALSE"
    )
    f["K"] = (
        f'=IF($A{r}="","",'
        f'IF(OR(_xlfn.ISFORMULA($A{r}),_xlfn.ISFORMULA($B{r})),"数式のまま貼り付けられています(値のみ貼り付けにしてください)",'
        f'IF(ISNA(MATCH($A{r},{WS_A},0)),"WSシートにないWS名です",'
        f'IF({order_check},"並び順が小さいタスクが下の行にあります(WS名→並び順で並べ替えてください)",'
        f'IF(AND($AB{r}="なし",$AA{r}="未定"),"WSの開始日が未入力です",'
        f'IF(AND(ISNUMBER($F{r}),ISNUMBER($G{r}),$G{r}<$F{r}),"終了日が開始日より前です",'
        f'IF(AND($E{r}="Yes",ISNUMBER($F{r}),ISNUMBER($AC{r})),'
        f'IF($F{r}<$AC{r},"前のタスクの終了日("&TEXT($AC{r}-1,"yyyy/mm/dd")&")以前です",""),"")))))))'
    )
    for col, formula in f.items():
        ws_t[f"{col}{r}"] = formula
    for col in ("C", "D", "E", "F", "G", "H", "J", "N"):
        ws_t[f"{col}{r}"].alignment = Alignment(horizontal="center")

add_list(ws_t, f"A2:A{T_LAST}", f"=WS!$A$2:$A${WS_LAST}")
add_list(ws_t, f"E2:E{T_LAST}", '"Yes,No"')
add_list(ws_t, f"N2:N{T_LAST}", '"' + ",".join(STATUSES) + '"')
add_list(ws_t, f"{AS1}2:{ASN}{T_LAST}", MEMBERS_LIST)
add_date_rule(ws_t, f"L2:M{T_LAST}")
add_int_rule(ws_t, f"C2:D{T_LAST}")
# WSごとに縞模様、未定・決定待ち・警告・ステータスを色分け
ws_t.conditional_formatting.add(
    f"A2:E{T_LAST}",
    FormulaRule(formula=['AND($A2<>"",ISNUMBER($AD2),ISODD($AD2))'], fill=PatternFill("solid", start_color="EAF1FB")),
)
ws_t.conditional_formatting.add(
    f"F2:G{T_LAST}",
    FormulaRule(formula=['F2="未定"'], font=Font(name=FONT, color="808080", italic=True)),
)
ws_t.conditional_formatting.add(
    f"J2:J{T_LAST}",
    FormulaRule(formula=['J2="決定待ち"'], font=Font(name=FONT, color="C55A11", bold=True)),
)
ws_t.conditional_formatting.add(
    f"K2:K{T_LAST}",
    FormulaRule(formula=['K2<>""'], fill=PatternFill("solid", start_color="F8CBAD"),
                font=Font(name=FONT, color="C00000", bold=True)),
)
for st, color in STATUS_FILLS.items():
    ws_t.conditional_formatting.add(
        f"N2:N{T_LAST}", FormulaRule(formula=[f'$N2="{st}"'], fill=PatternFill("solid", start_color=color)),
    )
ws_t.column_dimensions.group("Z", "AD", hidden=True, outline_level=1)
ws_t.auto_filter.ref = f"A1:{MEMO}{T_LAST}"

# サンプル: タスク生成で作ったのと同じ行 + WS固有のタスク + 入力値
SAMPLE_INPUTS = {
    # (WS名, タスク名): (開始日(手入力), 終了日(手入力), ステータス, [担当者], メモ)
    ("WS 秋季研修", "コンテンツ開発"): (None, None, "完了", ["田中 三郎"], "担当者だけ入力(日付は自動)"),
    ("WS 秋季研修", "場所決め"): (d("2026/09/28"), d("2026/10/11"), "完了", ["村上 四郎"], "開始日・終了日を手動で指定"),
    ("WS 秋季研修", "リハーサル物品購入"): (None, None, "着手中", ["鈴木 次郎"], None),
    ("WS 秋季研修", "リハーサル&広報撮影"): (d("2026/11/10"), None, None, ["佐藤 太郎", "鈴木 次郎"], "日付決定"),
    ("WS 秋季研修", "ワークショップ実施"): (
        d("2026/12/20"), None, None, ["佐藤 太郎", "鈴木 次郎", "田中 三郎", "村上 四郎"],
        "日付決定(広報稼働と重なるため警告が出る例)"),
    ("WS 秋季研修", "会場下見"): (None, None, None, ["村上 四郎"], "このWSだけのタスク(WS固有)"),
    ("Snow Dorm", "コンテンツ開発"): (None, None, "着手中", ["鈴木 次郎"], None),
}
SAMPLE_EXTRA = [("WS 秋季研修", "会場下見", 2, 3, "No")]  # WS固有のタスク(リハーサルと並列)


def generated_rows(ws_name):
    """タスク生成シートと同じ並び(標準並び順 → マスタの行の順)のタスク。"""

    ordered = sorted(enumerate(MASTERS), key=lambda x: (x[1][2], x[0]))
    return [(ws_name, name, order, days, yes) for _, (name, days, order, yes) in ordered]


if SAMPLE:
    rows = []
    for ws_name in ("WS 秋季研修", "Snow Dorm"):
        rows += generated_rows(ws_name)
        rows += [x for x in SAMPLE_EXTRA if x[0] == ws_name]
    rows.sort(key=lambda x: ([w[0] for w in SAMPLE_WS].index(x[0]), x[2]))  # WS名 → 並び順(安定ソート)
    for i, (ws_name, name, order, days, yes) in enumerate(rows, start=2):
        for col, v in zip("ABCDE", (ws_name, name, order, days, yes)):
            ws_t[f"{col}{i}"] = v
        inp = SAMPLE_INPUTS.get((ws_name, name))
        if inp:
            start, end, status, names, memo = inp
            ws_t[f"L{i}"], ws_t[f"M{i}"], ws_t[f"N{i}"], ws_t[f"{MEMO}{i}"] = start, end, status, memo
            for k, n in enumerate(names):
                ws_t.cell(row=i, column=15 + k, value=n)

# ---------------------------------------------------------------------------
# タスク生成(WSを選ぶと、タスクマスタからそのWSのタスクの行を作る。値のみ貼り付けでタスク一覧へ)
# ---------------------------------------------------------------------------
ws_g = wb.create_sheet("タスク生成")
ws_g.sheet_view.showGridLines = False
G_LAST = GEN_FIRST + N_MASTER - 1
gen_cols = [("WS名", 18), ("タスク名", 22), ("並び順", 8), ("所要日数", 9), ("日付決定有無", 11), ("", 2), ("マスタ行", 8)]
for i, (_, width) in enumerate(gen_cols, start=1):
    ws_g.column_dimensions[L(i)].width = width
ws_g["A1"] = "① WSを選ぶ"
ws_g["A2"] = "状態"
ws_g["A3"] = "手順"
for c in ("A1", "A2", "A3"):
    ws_g[c].font = base_font(bold=True)
ws_g["B1"].fill = INPUT_FILL
ws_g["B1"].border = BORDER
ws_g["B1"].font = base_font(bold=True)
ws_g["B1"].comment = Comment("WSシートに入力したWS名から選びます", "template")
add_list(ws_g, "B1", f"=WS!$A$2:$A${WS_LAST}")
gen_count = f'COUNTIF($A${GEN_FIRST}:$A${G_LAST},"?*")'
done_count = f"COUNTIF({TQ}!$A$2:$A${T_LAST},$B$1)"
ws_g["B2"] = (
    f'=IF($B$1="","B1 でWSを選んでください",IF(ISNA(MATCH($B$1,WS!$A$2:$A${WS_LAST},0)),"WSシートにないWS名です",'
    f'IF({done_count}>0,"このWSはすでにタスク一覧に "&{done_count}&" 件あります。二重に貼り付けないよう注意してください",'
    f'"未生成です。下の "&{gen_count}&" 行をコピーして、タスク一覧に値のみ貼り付けてください")))'
)
ws_g["B2"].font = base_font(bold=True, color="C00000")
steps = [
    "② 下の表の WS名〜日付決定有無(A〜E列)のタスクの行を選んでコピーします。",
    "③ 「タスク一覧」の一番下の空き行の A列(WS名)のセルを選び、「値のみ貼り付け」します"
    "(右クリック → 貼り付けのオプション → 値、または ホーム → 貼り付け → 値の貼り付け)。",
    "④ WS固有のタスクがあれば、その下の空き行に追加します。並び順が前後したら WS名 → 並び順 で並べ替えます。",
]
for i, text in enumerate(steps, start=3):
    ws_g[f"B{i}"] = text
    ws_g[f"B{i}"].font = base_font()
header(ws_g, gen_cols, [HDR_AUTO] * 5 + [NO_FILL, HDR_HELP], row=GEN_FIRST - 1)
ws_g.freeze_panes = None
body_style(ws_g, 1, 5, GEN_FIRST, G_LAST, AUTO_FILL)
body_style(ws_g, 7, 7, GEN_FIRST, G_LAST, AUTO_FILL, font_color="595959")
TM = lambda col: f"{TMQ}!${col}$1:${col}${M_LAST}"  # noqa: E731
for r in range(GEN_FIRST, G_LAST + 1):
    ws_g[f"G{r}"] = f"=IFERROR(MATCH(ROW()-{GEN_FIRST - 1},{TM('F')},0),\"\")"
    ws_g[f"A{r}"] = f'=IF(OR($B$1="",$G{r}=""),"",$B$1)'
    ws_g[f"B{r}"] = f'=IF($A{r}="","",INDEX({TM("A")},$G{r}))'
    ws_g[f"C{r}"] = f'=IF($A{r}="","",INDEX({TM("C")},$G{r}))'
    ws_g[f"D{r}"] = f'=IF($A{r}="","",INDEX({TM("B")},$G{r}))'
    ws_g[f"E{r}"] = f'=IF($A{r}="","",IF(INDEX({TM("D")},$G{r})="Yes","Yes","No"))'
    for col in "CDE":
        ws_g[f"{col}{r}"].alignment = Alignment(horizontal="center")
ws_g.column_dimensions.group("G", "G", hidden=True, outline_level=1)

# ---------------------------------------------------------------------------
# 個別タスク
# ---------------------------------------------------------------------------
ws_o = wb.create_sheet("個別タスク")
header(
    ws_o,
    [("WS名", 20), ("タスク名", 22), ("開始日", 12), ("終了日", 12), ("作業日数", 10), ("ステータス", 10)]
    + [(f"担当者{i}", 12) for i in range(1, N_ASSIGNEES + 1)]
    + [("メモ", 24)],
    HDR_INPUT,
    {
        "WS名": "空欄でよい(空欄 = その他)。WSに属さない作業(社内MTGなど)をここに直接書きます",
        "終了日": "空欄なら 開始日+作業日数−1",
        "ステータス": "未着手 / 着手中 / 完了(空欄は未着手)",
    },
)
lo = N_OTHER + 1
o_last = 7 + N_ASSIGNEES
body_style(ws_o, 1, o_last, 2, lo, INPUT_FILL, date_cols=("C", "D"))
for i, row in enumerate(
    [] if not SAMPLE else [
        [None, "社内MTG", d("2026/10/06"), None, 1, "完了", "佐藤 太郎", "鈴木 次郎", "田中 三郎"],
        [None, "資料作成", d("2026/10/13"), d("2026/10/16"), None, "未着手", "田中 三郎"],
    ],
    start=2,
):
    for j, v in enumerate(row, start=1):
        ws_o.cell(row=i, column=j, value=v)
add_list(ws_o, f"A2:A{lo}", f"=WS!$A$2:$A${WS_LAST}")
add_list(ws_o, f"F2:F{lo}", '"' + ",".join(STATUSES) + '"')
add_list(ws_o, f"G2:{L(6 + N_ASSIGNEES)}{lo}", MEMBERS_LIST)
add_date_rule(ws_o, f"C2:D{lo}")
add_int_rule(ws_o, f"E2:E{lo}")
for st, color in STATUS_FILLS.items():
    ws_o.conditional_formatting.add(
        f"F2:F{lo}", FormulaRule(formula=[f'$F2="{st}"'], fill=PatternFill("solid", start_color=color)),
    )
ws_o.auto_filter.ref = f"A1:{L(o_last)}{lo}"

# シートの並び: 説明 / タスク一覧 / WS / タスク生成 / タスクマスタ / メンバー / 個別タスク
order = ["説明", "タスク一覧", "WS", "タスク生成", "タスクマスタ", "メンバー", "個別タスク"]
wb._sheets = [wb[n] for n in order]
wb["タスク一覧"].sheet_properties.tabColor = "305496"
wb["タスク生成"].sheet_properties.tabColor = "548235"
for n in ("WS", "タスクマスタ", "メンバー", "個別タスク"):
    wb[n].sheet_properties.tabColor = "305496"
wb.active = 1
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT)
