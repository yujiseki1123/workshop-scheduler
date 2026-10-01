"""古い形式のスケジュール管理ブックから、入力した値を新しい形式のブックへ移す。

使い方:
    python -m tools.upgrade_workbook data/workshop_schedule.xlsx data/workshop_schedule_new.xlsx

新しいブックを tools/make_schedule_workbook.py で作り、古いブックの入力値を見出し名で対応させて書き写す。
数式の列(自動計算の列)は新しいブックの数式のまま。

- WS・タスクマスタ・メンバー・個別タスク: 行ごとに見出し名で書き写す。
- タスク一覧:
    - 旧形式(「日程入力」シートがあるブック。タスク一覧がマスタから数式で自動生成される形式)の場合:
      旧タスク一覧の計算結果(WS名・タスク名・並び順・日付決定有無・作業日数の標準値)を値として書き、
      日程入力の入力値(開始日・終了日・ステータス・担当者・メモ)を「WS名|タスク名」で対応させて書き写す。
      → 以後、マスタを変えてもこれらのWSのタスクは変わらない。
      旧タスク一覧は数式なので、古いブックは Excel で保存されている(計算結果が入っている)必要がある。
    - 新形式どうしの場合: 入力の列を見出し名で行ごとに書き写す。

作成後は Excel で一度開いて保存すること(数式の計算結果を保存するため)。
"""

import argparse
import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT_SHEETS = ("WS", "タスクマスタ", "メンバー", "個別タスク")
TASKS = "タスク一覧"
OLD_INPUT = "日程入力"
N_ASSIGNEES = 10


def _headers(ws):
    """見出し名(改行の前)→ 列番号。"""

    out = {}
    for cell in ws[1]:
        if cell.value:
            out.setdefault(str(cell.value).split("\n")[0].strip(), cell.column)
    return out


def _is_formula(v):
    return isinstance(v, str) and v.startswith("=")


def _input_columns(new_ws):
    """新しいブックで「入力する列」= 2行目が数式でない列。"""

    return {name: col for name, col in _headers(new_ws).items() if not _is_formula(new_ws.cell(2, col).value)}


def _text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _copy_by_header(o, n, notes, name):
    o_cols, n_cols = _headers(o), _input_columns(n)
    common = [h for h in n_cols if h in o_cols]
    for r in range(2, n.max_row + 1):
        for h in common:
            n.cell(r, n_cols[h]).value = None
    w = 2
    for r in range(2, o.max_row + 1):
        values = {h: o.cell(r, o_cols[h]).value for h in o_cols}
        if all(v in (None, "") for h, v in values.items() if h in n_cols):
            continue
        for h in common:
            n.cell(w, n_cols[h]).value = values[h]
        w += 1
    missing = [h for h in o_cols if h not in n_cols and not _is_formula(o.cell(2, o_cols[h]).value)
               and "(自動)" not in str(o.cell(1, o_cols[h]).value)]
    if missing:
        notes.append(f"[{name}] 新しい形式に無い列は書き写していません: {', '.join(missing)}")


def _tasks_from_old_format(old_values, new_ws, notes):
    """旧形式: 旧タスク一覧の計算結果 + 日程入力の入力値 → 新タスク一覧の値。"""

    t, inp = old_values[TASKS], old_values[OLD_INPUT]
    tc, ic, nc = _headers(t), _headers(inp), _input_columns(new_ws)

    # 旧マスタの標準作業日数(旧タスク一覧に「標準日数」が無い場合の予備)
    master_days = {}
    if "タスクマスタ" in old_values.sheetnames:
        m = old_values["タスクマスタ"]
        mc = _headers(m)
        for r in range(2, m.max_row + 1):
            name = _text(m.cell(r, mc.get("タスク名", 1)).value)
            if name and "標準作業日数" in mc:
                master_days[name] = m.cell(r, mc["標準作業日数"]).value

    # 日程入力: 「WS名|タスク名」→ 行番号(A・B列は数式なので計算結果を使う)
    inputs = {}
    for r in range(2, inp.max_row + 1):
        ws_name, task = _text(inp.cell(r, ic["WS名"]).value), _text(inp.cell(r, ic["タスク名"]).value)
        if ws_name and task:
            inputs[(ws_name, task)] = r
        elif any(inp.cell(r, ic[h]).value not in (None, "") for h in ("開始日", "終了日", "ステータス", "メモ") if h in ic):
            notes.append(f"[日程入力] {r}行目: 対応するタスクが無い入力は移していません")

    rows = []
    for r in range(2, t.max_row + 1):
        ws_name = _text(t.cell(r, tc["WS名"]).value)
        if not ws_name:
            continue
        task = _text(t.cell(r, tc["タスク名"]).value)
        days = t.cell(r, tc["標準日数"]).value if "標準日数" in tc else None
        if days in (None, ""):
            days = master_days.get(task)
        rows.append({
            "WS名": ws_name,
            "タスク名": task,
            "並び順": t.cell(r, tc["並び順"]).value,
            "所要日数": days,
            "日付決定有無": "Yes" if _text(t.cell(r, tc["日付決定有無"]).value) == "Yes" else "No",
            "_input": inputs.get((ws_name, task)),
        })
    if not rows and any(_text(c.value) for c in old_values["WS"]["A"][1:]):
        raise SystemExit(
            "古いブックのタスク一覧に計算結果がありません。Excel(ブラウザ版可)で古いブックを開いて保存してから実行してください。"
        )

    copy_from_input = [("開始日", "開始日(手入力)"), ("終了日", "終了日(手入力)"), ("ステータス", "ステータス"),
                       ("メモ", "メモ")] + [(f"担当者{i}", f"担当者{i}") for i in range(1, N_ASSIGNEES + 1)]
    for w, row in enumerate(rows, start=2):
        for h in ("WS名", "タスク名", "並び順", "所要日数", "日付決定有無"):
            new_ws.cell(w, nc[h]).value = row[h]
        src = row["_input"]
        if src:
            for old_h, new_h in copy_from_input:
                if old_h in ic and new_h in nc:
                    v = inp.cell(src, ic[old_h]).value
                    if v not in (None, ""):
                        new_ws.cell(w, nc[new_h]).value = v
    notes.append(f"[タスク一覧] 旧形式から {len(rows)} 件のタスクを値として移しました(以後マスタを変えても変わりません)")


def upgrade(old_path, new_path):
    subprocess.run(
        [sys.executable, "-m", "tools.make_schedule_workbook", str(new_path)], cwd=ROOT, check=True,
        capture_output=True,
    )
    old = load_workbook(old_path, data_only=True)
    new = load_workbook(new_path)
    notes = []

    for name in INPUT_SHEETS:
        if name in old.sheetnames:
            _copy_by_header(old[name], new[name], notes, name)

    if OLD_INPUT in old.sheetnames:
        _tasks_from_old_format(old, new[TASKS], notes)
    elif TASKS in old.sheetnames:
        _copy_by_header(old[TASKS], new[TASKS], notes, TASKS)

    new.calculation.fullCalcOnLoad = True
    new.save(new_path)
    return notes


def main():
    parser = argparse.ArgumentParser(description="古い形式のブックから入力値を新しい形式へ移す")
    parser.add_argument("old")
    parser.add_argument("new")
    args = parser.parse_args()
    notes = upgrade(args.old, args.new)
    print(f"作成しました: {args.new}(Excelで一度開いて保存してください)")
    for n in notes:
        print("  - " + n)


if __name__ == "__main__":
    main()
