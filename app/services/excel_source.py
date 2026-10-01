"""スケジュール管理ブック(Excel)を読み取る層。

Excel が唯一の正(Single Source of Truth)。Webアプリは読むだけで書き換えない。
ブックの形式は tools/make_schedule_workbook.py で作るもの。

読むシート:
    WS         : 行の一覧(WS名・種別・並び順)
    タスク一覧 : 数式で計算済みのタスク(開始日・終了日・担当者・警告)
    個別タスク : マスタから自動で作らない作業(社内MTGなど)
    メンバー   : 担当者の並び順
    タスクマスタ: /api/task-masters 用

注意: タスク一覧は数式なので、Excel が保存した「計算結果」を読む(openpyxl の data_only)。
openpyxl など Excel 以外で作っただけのファイルには計算結果が無いため、一度 Excel で
開いて保存してから読ませる必要がある(その場合は ExcelNotCalculatedError)。

ファイルの更新日時とサイズが変わったときだけ読み直す(読み込み結果をキャッシュ)。
将来 OneDrive から読む場合も、ファイルをダウンロードしてこの層に渡せばよい。
"""

import os
import threading
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

SHEET_WS = "WS"
SHEET_TASKS = "タスク一覧"
SHEET_OTHER = "個別タスク"
SHEET_MEMBERS = "メンバー"
SHEET_MASTERS = "タスクマスタ"
UNDECIDED = "未定"
CATEGORY_WORKSHOP = "ワークショップ"
CATEGORY_OTHER = "その他"
ASSIGNEE_SEPARATOR = "、"
STATUSES = ("未着手", "着手中", "完了")
DEFAULT_STATUS = "未着手"
MAX_ASSIGNEES = 10


class ExcelSourceError(Exception):
    """ブックを読めない・形式が違うときのエラー(画面にそのまま表示できる文言)。"""


class ExcelNotCalculatedError(ExcelSourceError):
    pass


@dataclass
class ScheduleBook:
    projects: list = field(default_factory=list)
    tasks: list = field(default_factory=list)
    members: list = field(default_factory=list)
    task_masters: list = field(default_factory=list)
    warnings: list = field(default_factory=list)  # 読み取り時に見つかった問題(表示は続ける)
    path: str = ""
    modified_at: str = ""


# ---------------------------------------------------------------------------
# セル値の変換
# ---------------------------------------------------------------------------


def _text(v):
    if v is None:
        return None
    s = str(v).strip()
    return s or None


def _int(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return int(v)
    try:
        return int(str(v).strip())
    except ValueError:
        return None


def _date(v):
    """日付セル → date。「未定」や空欄は None。"""

    if v is None:
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # 書式が日付でない場合、Excelのシリアル値のまま入っていることがある
        return from_excel(v).date()
    s = str(v).strip()
    if not s or s == UNDECIDED:
        return None
    for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    return None


def _status(v):
    s = _text(v)
    return s if s in STATUSES else DEFAULT_STATUS


def _rows(ws):
    """シートの行を {見出し: 値} で返す。見出しは改行の前(「WS名\\n(自動)」→「WS名」)で照合する。"""

    it = ws.iter_rows(values_only=True)
    try:
        header = next(it)
    except StopIteration:
        return
    names = [(_text(h) or "").split("\n")[0] for h in header]
    for row_no, cells in enumerate(it, start=2):
        yield row_no, {n: cells[i] if i < len(cells) else None for i, n in enumerate(names) if n}


def _require(wb, name):
    if name not in wb.sheetnames:
        raise ExcelSourceError(f"シート「{name}」がありません(スケジュール管理ブックの形式か確認してください)")
    return wb[name]


# ---------------------------------------------------------------------------
# 読み取り
# ---------------------------------------------------------------------------


def parse_workbook(path) -> ScheduleBook:
    if not os.path.exists(path):
        raise ExcelSourceError(f"Excelファイルが見つかりません: {path}")
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception as e:  # 壊れたファイル・Excel以外のファイル
        raise ExcelSourceError(f"Excelファイルを開けません: {path}({e})") from e

    book = ScheduleBook(
        path=path,
        modified_at=datetime.fromtimestamp(os.path.getmtime(path)).isoformat(timespec="seconds"),
    )
    try:
        _read_members(wb, book)
        _read_masters(wb, book)
        _read_projects(wb, book)
        _read_tasks(wb, book)
        _read_other_tasks(wb, book)
    finally:
        wb.close()

    order = {p["project_id"]: i for i, p in enumerate(book.projects)}
    book.tasks.sort(key=lambda t: (order[t["project_id"]], t["task_order"], t["_row"]))
    for i, t in enumerate(book.tasks, start=1):
        t["task_id"] = i
        t.pop("_row")
    return book


def _read_members(wb, book):
    if SHEET_MEMBERS not in wb.sheetnames:
        return
    for _, r in _rows(wb[SHEET_MEMBERS]):
        name = _text(r.get("表示名")) or " ".join(filter(None, [_text(r.get("姓")), _text(r.get("名"))]))
        if not name:
            continue
        book.members.append(
            {
                "member_id": len(book.members) + 1,
                "display_name": name,
                "email": _text(r.get("メールアドレス")),
                "is_active": _text(r.get("状態")) != "無効",
            }
        )


def _read_masters(wb, book):
    if SHEET_MASTERS not in wb.sheetnames:
        return
    for _, r in _rows(wb[SHEET_MASTERS]):
        name = _text(r.get("タスク名"))
        if not name:
            continue
        book.task_masters.append(
            {
                "task_master_id": len(book.task_masters) + 1,
                "task_name": name,
                "default_work_days": _int(r.get("標準作業日数")),
                "default_task_order": _int(r.get("標準並び順")),
                "requires_date_decision": _text(r.get("日付決定有無")) == "Yes",
                "auto_generate": _text(r.get("WS自動生成")) != "対象外",
            }
        )


def _read_projects(wb, book):
    rows = []
    for row_no, r in _rows(_require(wb, SHEET_WS)):
        name = _text(r.get("WS名"))
        if not name:
            continue
        if any(p["project_name"] == name for p in rows):
            book.warnings.append(f"[WS] {row_no}行目: WS名「{name}」が重複しています(2件目以降は無視)")
            continue
        category = _text(r.get("種別")) or CATEGORY_WORKSHOP
        rows.append(
            {
                "project_id": len(rows) + 1,
                "project_name": name,
                "project_category": category,
                "sort_order": _int(r.get("並び順")) if _int(r.get("並び順")) is not None else 9999,
                "start_date": _date(r.get("開始日")),
                "status": _text(r.get("ステータス")),
            }
        )
    rows.sort(key=lambda p: (p["sort_order"], p["project_id"]))
    book.projects = rows


def _member_refs(book, names):
    """担当者名 → {member_id, display_name}。メンバーシートに無い名前も表示はする。"""

    by_name = {m["display_name"]: m for m in book.members}
    out = []
    for n in names:
        m = by_name.get(n)
        if m is None:
            m = {"member_id": 100000 + len(by_name), "display_name": n}
            by_name[n] = m
            book.members.append({**m, "email": None, "is_active": True, "unregistered": True})
        out.append({"member_id": m["member_id"], "display_name": m["display_name"]})
    return out


def _task(book, project, row_no, name, order, days, start, end, names, **extra):
    t = {
        "_row": row_no,
        "project_id": project["project_id"],
        "task_name": name,
        "task_order": order if order is not None else 9999,
        "work_days": days,
        "start_date": start.isoformat() if start else None,
        "end_date": end.isoformat() if end else None,
        "is_undecided": start is None,
        "assignees": _member_refs(book, names),
    }
    t.update(extra)
    return t


def _read_tasks(wb, book):
    projects = {p["project_name"]: p for p in book.projects}
    ws = _require(wb, SHEET_TASKS)
    workshops = [p for p in book.projects if p["project_category"] == CATEGORY_WORKSHOP]
    has_masters = any(m["auto_generate"] for m in book.task_masters)
    found = 0
    for row_no, r in _rows(ws):
        ws_name = _text(r.get("WS名"))
        if not ws_name:
            continue
        found += 1
        project = projects.get(ws_name)
        if project is None:
            book.warnings.append(f"[タスク一覧] {row_no}行目: WS「{ws_name}」がWSシートにありません")
            continue
        start = _date(r.get("開始日"))
        days = _int(r.get("作業日数")) or 1
        end = _date(r.get("終了日")) or (start + timedelta(days=days - 1) if start else None)
        names = [n.strip() for n in (_text(r.get("担当者")) or "").split(ASSIGNEE_SEPARATOR) if n.strip()]
        book.tasks.append(
            _task(
                book, project, row_no, _text(r.get("タスク名")), _int(r.get("並び順")), days, start, end, names,
                requires_date_decision=_text(r.get("日付決定有無")) == "Yes",
                start_mode=_text(r.get("開始日の決め方")),
                status=_status(r.get("ステータス")),
                warning=_text(r.get("警告")),
                memo=_text(r.get("メモ")),
                source="タスク一覧",
            )
        )
    # ワークショップがあるのにタスク一覧が空 = 数式の計算結果がファイルに保存されていない
    if workshops and has_masters and found == 0:
        raise ExcelNotCalculatedError(
            "タスク一覧に計算結果がありません。Excel(ブラウザ版可)でファイルを一度開いて保存してから再読み込みしてください。"
        )


def _read_other_tasks(wb, book):
    if SHEET_OTHER not in wb.sheetnames:
        return
    projects = {p["project_name"]: p for p in book.projects}
    for row_no, r in _rows(wb[SHEET_OTHER]):
        ws_name, name = _text(r.get("WS名")), _text(r.get("タスク名"))
        if not ws_name and not name:
            continue
        where = f"[個別タスク] {row_no}行目"
        project = projects.get(ws_name)
        if project is None or not name:
            book.warnings.append(f"{where}: WS名・タスク名を確認してください(WS「{ws_name}」)")
            continue
        start = _date(r.get("開始日"))
        end = _date(r.get("終了日"))
        days = _int(r.get("作業日数")) or 1
        warning = None
        if start and end:
            days = (end - start).days + 1
            if end < start:
                warning = "終了日が開始日より前です"
        elif start:
            end = start + timedelta(days=days - 1)
        else:
            end = None  # 開始日が無ければ未定
        names = [n for n in (_text(r.get(f"担当者{i}")) for i in range(1, MAX_ASSIGNEES + 1)) if n]
        book.tasks.append(
            _task(
                book, project, 10000 + row_no, name, None, days, start, end, names,
                requires_date_decision=False,
                start_mode="手動",
                status=_status(r.get("ステータス")),
                warning=warning,
                memo=_text(r.get("メモ")),
                source="個別タスク",
            )
        )


# ---------------------------------------------------------------------------
# キャッシュ付きの読み込み
# ---------------------------------------------------------------------------

_cache = {"key": None, "book": None}
_lock = threading.Lock()


def load_book(path) -> ScheduleBook:
    """ファイルが変わっていなければ前回の読み込み結果を返す。"""

    try:
        st = os.stat(path)
        key = (os.path.abspath(path), st.st_mtime_ns, st.st_size)
    except FileNotFoundError:
        raise ExcelSourceError(f"Excelファイルが見つかりません: {path}") from None
    with _lock:
        if _cache["key"] == key:
            return _cache["book"]
        book = parse_workbook(path)
        _cache.update(key=key, book=book)
        return book


def clear_cache():
    with _lock:
        _cache.update(key=None, book=None)
