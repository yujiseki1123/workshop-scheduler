"""Excel(スケジュール管理ブック)を読み取る層のテスト。フィクスチャは conftest.py。"""

import shutil

import pytest

from app.services.excel_source import ExcelNotCalculatedError, ExcelSourceError, parse_workbook
from tests.conftest import TASK_ROWS, TEN, d, fill


# ---------------------------------------------------------------------------
# 読み取り
# ---------------------------------------------------------------------------


def test_projects_sorted_by_sort_order(book_path):
    book = parse_workbook(str(book_path))
    assert [p["project_name"] for p in book.projects] == ["WS 秋", "Snow Dorm", "その他"]
    assert book.projects[2]["project_category"] == "その他"


def test_tasks_dates_and_undecided(book_path):
    book = parse_workbook(str(book_path))
    tasks = {(t["task_name"], t["project_id"]): t for t in book.tasks}
    aki, snow = book.projects[0]["project_id"], book.projects[1]["project_id"]

    lp = tasks[("LP作成", aki)]
    assert (lp["start_date"], lp["end_date"], lp["work_days"]) == ("2026-11-11", "2026-11-24", 14)
    assert not lp["is_undecided"] and lp["start_mode"] == "自動"

    reh = tasks[("リハーサル&広報撮影", snow)]
    assert reh["start_date"] is None and reh["end_date"] is None and reh["is_undecided"]
    assert reh["requires_date_decision"] and reh["start_mode"] == "決定待ち"

    ws_run = tasks[("ワークショップ実施", aki)]
    assert "以前です" in ws_run["warning"]


def test_assignees_split_and_unregistered_names(book_path):
    book = parse_workbook(str(book_path))
    t = next(t for t in book.tasks if t["task_name"] == "リハーサル&広報撮影" and t["start_date"])
    assert [a["display_name"] for a in t["assignees"]] == ["佐藤 太郎", "鈴木 次郎"]
    assert [a["member_id"] for a in t["assignees"]] == [1, 2]  # メンバーシートの順
    # メンバーシートに無い名前も表示する(IDは登録済みの後ろ)
    outsider = next(t for t in book.tasks if t["assignees"] and t["assignees"][0]["display_name"] == "外部 花子")
    assert outsider["assignees"][0]["member_id"] > 3


def test_other_tasks_and_warnings(book_path):
    book = parse_workbook(str(book_path))
    mtg = next(t for t in book.tasks if t["task_name"] == "社内MTG")
    assert (mtg["start_date"], mtg["end_date"], mtg["source"]) == ("2026-10-06", "2026-10-06", "個別タスク")
    assert any("存在しないWS" in w for w in book.warnings)


def test_status_defaults_to_not_started(book_path):
    book = parse_workbook(str(book_path))
    by_name = {(t["task_name"], t["start_date"]): t["status"] for t in book.tasks}
    assert by_name[("コンテンツ開発", "2026-10-01")] == "完了"
    assert by_name[("リハーサル&広報撮影", "2026-11-10")] == "着手中"
    assert by_name[("LP作成", None)] == "未着手"  # 空欄は未着手
    assert by_name[("社内MTG", "2026-10-06")] == "完了"


def test_other_task_end_date_and_ten_assignees(book_path):
    book = parse_workbook(str(book_path))
    t = next(t for t in book.tasks if t["task_name"] == "合宿準備")
    assert (t["start_date"], t["end_date"], t["work_days"], t["status"]) == ("2026-10-13", "2026-10-16", 4, "着手中")
    assert [a["display_name"] for a in t["assignees"]] == TEN


def test_task_order_follows_projects(book_path):
    book = parse_workbook(str(book_path))
    names = [(t["project_id"], t["task_order"]) for t in book.tasks]
    assert names == sorted(names, key=lambda x: ([p["project_id"] for p in book.projects].index(x[0]), x[1]))
    assert [t["task_id"] for t in book.tasks] == list(range(1, len(book.tasks) + 1))


def test_masters_are_read(book_path):
    book = parse_workbook(str(book_path))
    assert len(book.task_masters) == 11
    wsx = next(m for m in book.task_masters if m["task_name"] == "ワークショップ実施")
    assert wsx["requires_date_decision"] and wsx["auto_generate"]


def test_empty_book_has_no_data(empty_book):
    book = parse_workbook(str(empty_book))
    assert book.projects == [] and book.tasks == [] and book.members == []


def test_not_calculated_book_is_reported(empty_book, tmp_path):
    """タスクの行はあるのに自動計算の列が空(openpyxl で作っただけで Excel で保存されていない)。"""

    path = tmp_path / "raw.xlsx"
    shutil.copy(empty_book, path)
    fill(path, ws_rows=[["WS 秋", "ワークショップ", 1, d("2026/10/01"), None]],
         task_rows=[["WS 秋", "コンテンツ開発", 1, 28, "No"]], task_columns=["WS名", "タスク名", "並び順", "所要日数", "日付決定有無"])
    with pytest.raises(ExcelNotCalculatedError):
        parse_workbook(str(path))


def test_workshop_without_tasks_is_warned(empty_book, tmp_path):
    """タスクをまだ生成していないWSは、エラーにせず注意として出す。"""

    path = tmp_path / "new_ws.xlsx"
    shutil.copy(empty_book, path)
    fill(path, ws_rows=[["WS 秋", "ワークショップ", 1, d("2026/10/01"), None], ["その他", "その他", 99, None, None]])
    book = parse_workbook(str(path))
    assert book.tasks == []
    assert any("「WS 秋」のタスクがまだありません" in w for w in book.warnings)
    assert not any("その他" in w for w in book.warnings)


def test_row_without_formulas_is_warned(book_path):
    """途中に挿入した行など、自動計算の数式が無い行は注意に出す(表示は未定)。"""

    fill(book_path, task_rows=TASK_ROWS + [["", "WS 秋", "挿入した行", 2, None, "No"]],
         ws_rows=[["Snow Dorm", "ワークショップ", 2, d("2026/11/02"), "企画中"],
                  ["WS 秋", "ワークショップ", 1, d("2026/10/01"), "準備中"], ["その他", "その他", 99, None, None]])
    book = parse_workbook(str(book_path))
    assert any("9行目: 自動計算の数式が入っていません" in w for w in book.warnings)
    t = next(t for t in book.tasks if t["task_name"] == "挿入した行")
    assert t["is_undecided"]


def test_old_format_book_still_reads(tmp_path):
    """旧形式(日程入力あり)のブックも読める(移行前のファイルを開いたとき用)。"""

    from pathlib import Path

    from openpyxl import load_workbook
    old = Path(__file__).resolve().parents[1] / "data" / "workshop_schedule.xlsx"
    if not old.exists() or "日程入力" not in load_workbook(old, read_only=True).sheetnames:
        pytest.skip("旧形式のブックが無い")
    book = parse_workbook(str(old))
    assert book.tasks


def test_missing_file(tmp_path):
    with pytest.raises(ExcelSourceError):
        parse_workbook(str(tmp_path / "none.xlsx"))
